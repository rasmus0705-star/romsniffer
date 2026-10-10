"""
rom_overrides.py — facitlisten (fejlliste.xlsx) til RomSniffer. Samme princip som på BeerSniffer:
dine manuelle rettelser gælder FØR matchning, og facit vinder.

fejlliste.xlsx har to ark:
  "Fejlliste"  én række pr. vare, du har rettet:  Butik | Navn | URL | Brand | Volumen (cl) | ABV (%) | Status
               Status: tom = rettet, SAET = sæt/kalender (slås aldrig sammen med andre), IGNORERET = fjernes helt
  "Par"        beslutninger om to varer:           Navn A | Butik A | URL A | Navn B | Butik B | URL B | Beslutning | Note
               Beslutning: sammen (samme flaske) eller adskilt (forskellige flasker)
Filen kan også rettes direkte i Excel. Varer identificeres på URL.

fejl_snapshot.json skrives af hvert build og er det, start_Fejlrettelse.bat viser (hvad mangler, hvilke par skal
gennemgås). Den er diagnose, ikke facit, og skal ikke committes.
"""
import json
import os
import re
from datetime import datetime

from rom_text import clean_text, plausible_abv

FEJLLISTE = "fejlliste.xlsx"
SNAPSHOT = "fejl_snapshot.json"
SHEET_ITEMS = "Fejlliste"
SHEET_PAIRS = "Par"
ITEM_COLS = ["Butik", "Navn", "URL", "Brand", "Volumen (cl)", "ABV (%)", "Status"]
PAIR_COLS = ["Navn A", "Butik A", "URL A", "Navn B", "Butik B", "URL B", "Beslutning", "Note"]
LOCK_MSG = "fejlliste.xlsx er åben i Excel (eller låst). Luk den og prøv igen."


def pair_key(url_a, url_b):
    return tuple(sorted((url_a, url_b)))


def _num(v):
    """Tal fra en celle: accepterer 40,5 / 40.5 / '40,5 %' / tom."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^0-9,.\-]", "", str(v)).replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _s(v):
    return "" if v is None else str(v).strip()


def empty_facit():
    return {"items": {}, "pairs": {}, "pair_info": {}}


def load_fejlliste(path=FEJLLISTE):
    """Returnerer {"items": {url: {...}}, "pairs": {(url_a, url_b): "sammen"|"adskilt"}, "pair_info": {...}}."""
    facit = empty_facit()
    if not os.path.isfile(path):
        return facit
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, data_only=True)
    except PermissionError:
        raise SystemExit(LOCK_MSG)
    if SHEET_ITEMS in wb.sheetnames:
        rows = list(wb[SHEET_ITEMS].iter_rows(values_only=True))
        for r in rows[1:]:
            r = list(r) + [None] * (len(ITEM_COLS) - len(r))
            butik, navn, url, brand, vol, abv, status = r[:7]
            if not _s(url):
                continue
            facit["items"][_s(url)] = {
                "butik": _s(butik), "navn": _s(navn), "brand": _s(brand) or None,
                "volume": _num(vol), "abv": _num(abv), "status": _s(status).upper(),
            }
    if SHEET_PAIRS in wb.sheetnames:
        rows = list(wb[SHEET_PAIRS].iter_rows(values_only=True))
        for r in rows[1:]:
            r = list(r) + [None] * (len(PAIR_COLS) - len(r))
            navn_a, butik_a, url_a, navn_b, butik_b, url_b, beslutning, note = r[:8]
            dec = _s(beslutning).lower()
            if _s(url_a) and _s(url_b) and dec in ("sammen", "adskilt"):
                k = pair_key(_s(url_a), _s(url_b))
                facit["pairs"][k] = dec
                facit["pair_info"][k] = {"navn_a": _s(navn_a), "butik_a": _s(butik_a), "url_a": _s(url_a),
                                         "navn_b": _s(navn_b), "butik_b": _s(butik_b), "url_b": _s(url_b), "note": _s(note)}
    return facit


def save_fejlliste(facit, path=FEJLLISTE):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_ITEMS
    ws.append(ITEM_COLS)
    for url, f in sorted(facit["items"].items(), key=lambda kv: (kv[1].get("butik", ""), kv[1].get("navn", ""))):
        ws.append([f.get("butik", ""), f.get("navn", ""), url, f.get("brand") or "", f.get("volume"), f.get("abv"), f.get("status", "")])
    wp = wb.create_sheet(SHEET_PAIRS)
    wp.append(PAIR_COLS)
    for k, dec in facit["pairs"].items():
        i = facit["pair_info"].get(k, {})
        wp.append([i.get("navn_a", ""), i.get("butik_a", ""), i.get("url_a", k[0]), i.get("navn_b", ""),
                   i.get("butik_b", ""), i.get("url_b", k[1]), dec, i.get("note", "")])
    head = PatternFill("solid", fgColor="F0B832")
    for sheet, widths in ((ws, [14, 50, 40, 20, 12, 10, 12]), (wp, [38, 12, 34, 38, 12, 34, 12, 30])):
        for c in sheet[1]:
            c.font = Font(bold=True)
            c.fill = head
        for i, w in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(i)].width = w
        sheet.freeze_panes = "A2"
    try:
        wb.save(path)
    except PermissionError:
        raise RuntimeError(LOCK_MSG)


def apply_overrides(item, facit):
    """Læg dine rettelser ind over en vare. Kaldes FØR matchning."""
    f = facit["items"].get(item.get("url"))
    if not f:
        return item
    if f.get("brand"):
        item["brand"] = f["brand"]
    if f.get("volume") is not None:
        item["volume_cl"] = f["volume"]
    if f.get("abv") is not None:
        item["abv"] = f["abv"]
    status = (f.get("status") or "").upper()
    if status == "IGNORERET":
        item["_skip"] = True
    elif status == "SAET":
        item["_solo"] = True
    item["_facit"] = True
    return item


# ───────────────────────────── diagnose (fejl_snapshot.json) ─────────────────────────────
_NOT_RUM = re.compile(r"\b(gin|vodka|whisky|whiskey|tequila|likør|likor|liqueur|cognac|brandy|champagne|prosecco|gavekort|glas)\b", re.I)


def _known_brands():
    try:
        from rom_parser import KNOWN_BRANDS
        return list(KNOWN_BRANDS)
    except Exception:
        return []


def _guess_brand(name, brands):
    low = name.lower()
    for b in sorted(brands, key=len, reverse=True):
        bl = b.lower()
        if low.startswith(bl + " ") or low.startswith(bl + " -") or low.startswith(bl + " –"):
            return b, "HOEJ"
    for b in sorted(brands, key=len, reverse=True):
        if len(b) >= 5 and b.lower() in low:
            return b, "LAV"
    return None, None


def write_snapshot(items, groups, facit, path=SNAPSHOT, max_pairs=300):
    import rom_matching as M

    def key(it):
        return (it.get("shop_name"), it.get("url") or it.get("name"))

    gid, members = {}, {}
    for i, g in enumerate(groups):
        members[i] = g
        for p in g:
            gid[key(p)] = i

    rows = []
    for it in items:
        name = clean_text(it.get("name"))
        suspicious = it.get("abv") is not None and plausible_abv({"abv": it.get("abv"), "name": name}) is None
        rows.append({
            "url": it.get("url"), "butik": it.get("shop_name"), "navn": name, "brand": it.get("brand"),
            "volume": it.get("volume_cl"), "abv": it.get("abv"), "alder": it.get("age"), "pris": it.get("price"),
            "gruppe": gid.get(key(it)), "abv_mistaenkelig": suspicious,
        })

    brands = sorted({r["brand"] for r in rows if r["brand"]} | set(_known_brands()), key=str.lower)

    forslag = []
    for r in rows:
        if r["brand"] or not r["url"]:
            continue
        b, conf = _guess_brand(r["navn"], brands)
        if b:
            forslag.append({"url": r["url"], "navn": r["navn"], "gammelt": "", "forslag": b, "sikkerhed": conf})

    match = []
    decided = facit.get("pairs", {})
    # 1) OPSPLITTET: lignende navne i forskellige butikker, som ikke er samlet
    toks = [set(M.clean_name(it.get("name", "")).split()) for it in items]
    cands = []
    for i, a in enumerate(items):
        if not toks[i]:
            continue
        for j in range(i + 1, len(items)):
            b = items[j]
            if a.get("shop_name") == b.get("shop_name") or gid.get(key(a)) == gid.get(key(b)):
                continue
            inter = toks[i] & toks[j]
            if len(inter) < 2:
                continue
            sim = len(inter) / len(toks[i] | toks[j])
            if sim >= 0.6:
                cands.append((sim, a, b))
    cands.sort(key=lambda x: -x[0])
    for sim, a, b in cands:
        if len(match) >= max_pairs:
            break
        if pair_key(a.get("url"), b.get("url")) in decided:
            continue
        res = M.try_match(a, b)
        reason = (res.get("reason") or "").replace("\u274c ", "")
        if "volume-gate: Volume mismatch" in reason:
            continue                      # to kendte, forskellige størrelser (35 cl mod 70 cl): ikke en fejl
        if not reason:
            def mates(it):
                others = [p for p in members.get(gid.get(key(it)), []) if key(p) != key(it)]
                return ", ".join("%s: %s" % (p.get("shop_name"), clean_text(p.get("name"))) for p in others) or "står alene"
            reason = "matcher, men ligger i forskellige grupper (A sammen med: %s; B sammen med: %s)" % (mates(a), mates(b))
        match.append({"type": "OPSPLITTET", "alvor": "HØJ" if sim >= 0.85 else "MELLEM",
                      "ol_a": clean_text(a.get("name")), "butik_a": a.get("shop_name"), "url_a": a.get("url"),
                      "ol_b": clean_text(b.get("name")), "butik_b": b.get("shop_name"), "url_b": b.get("url"),
                      "detaljer": "%d%% ens navn. %s" % (sim * 100, reason)})

    # 2) PRIS-SPREDNING: samlede grupper med stor prisforskel (kan være to forskellige flasker)
    for g in groups:
        if len({p.get("shop_name") for p in g}) < 2:
            continue
        priced = [p for p in g if p.get("price")]
        if len(priced) < 2:
            continue
        lo = min(priced, key=lambda p: float(p["price"]))
        hi = max(priced, key=lambda p: float(p["price"]))
        ratio = float(hi["price"]) / float(lo["price"]) if float(lo["price"]) > 0 else 0
        if ratio >= 1.5 and pair_key(lo.get("url"), hi.get("url")) not in decided:
            match.append({"type": "PRIS-SPREDNING", "alvor": "HØJ" if ratio >= 1.75 else "MELLEM",
                          "ol_a": clean_text(lo.get("name")), "butik_a": lo.get("shop_name"), "url_a": lo.get("url"),
                          "ol_b": clean_text(hi.get("name")), "butik_b": hi.get("shop_name"), "url_b": hi.get("url"),
                          "detaljer": "%.0f kr mod %.0f kr (%.2fx), volumen %s mod %s cl" % (
                              float(lo["price"]), float(hi["price"]), ratio, lo.get("volume_cl"), hi.get("volume_cl"))})

    # 3) IKKE-ROM?: navne der ligner andre spirituosa
    for it in items:
        name = clean_text(it.get("name"))
        m = _NOT_RUM.search(name)
        if m and not it.get("_facit"):
            match.append({"type": "IKKE-ROM?", "alvor": "MELLEM", "ol_a": name, "butik_a": it.get("shop_name"),
                          "url_a": it.get("url"), "ol_b": "", "butik_b": "", "url_b": "",
                          "detaljer": "navnet indeholder \"%s\"" % m.group(1)})

    data = {"tidspunkt": datetime.now().strftime("%d-%m-%Y %H:%M"), "items": rows, "forslag": forslag,
            "match": match, "brands": brands}
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)
    n_susp = sum(1 for r in rows if r["abv_mistaenkelig"])
    print(f"📝 Fejl-snapshot: {len(rows)} varer, {len(match)} par/problemer at gennemgå, {len(forslag)} brand-forslag, {n_susp} mistænkelige ABV → {path}")
    return data
